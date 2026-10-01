from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows

from modules.orientation import restore_original_orientation


def corrected_table(normalized_df, orientation, original_df):
    return restore_original_orientation(normalized_df, orientation, original_df)


def _replace_sheet(ws, df: pd.DataFrame) -> None:
    ws.delete_rows(1, ws.max_row)
    for row in dataframe_to_rows(df, index=False, header=True):
        ws.append(list(row))


def build_corrected_workbook_bytes(
    original_bytes: bytes,
    filename: str,
    selected_sheet: str | None,
    original_df: pd.DataFrame,
    corrected_normalized_df: pd.DataFrame,
    orientation,
) -> bytes:
    corrected = corrected_table(corrected_normalized_df, orientation, original_df)
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    output = BytesIO()

    if ext == "xlsx":
        wb = load_workbook(BytesIO(original_bytes))
        sheet = selected_sheet or wb.sheetnames[0]
        ws = wb[sheet]
        _replace_sheet(ws, corrected)
        wb.save(output)
        return output.getvalue()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        corrected.to_excel(writer, sheet_name="Corrected_Data", index=False)
    return output.getvalue()


def build_audit_report_bytes(
    unresolved: pd.DataFrame,
    resolved: pd.DataFrame,
    audit_log: list[dict],
    lifecycle_summary: pd.DataFrame,
    fill_candidates: pd.DataFrame,
    summary: dict,
    coverage_candidates: pd.DataFrame | None = None,
) -> bytes:
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame([summary]).to_excel(writer, sheet_name="Summary", index=False)
        unresolved.to_excel(writer, sheet_name="Unresolved", index=False)
        resolved.to_excel(writer, sheet_name="Resolved", index=False)
        pd.DataFrame(audit_log or []).to_excel(writer, sheet_name="Change_Log", index=False)
        lifecycle_summary.to_excel(writer, sheet_name="Lifecycle", index=False)
        fill_candidates.to_excel(writer, sheet_name="Safe_Fill_Preview", index=False)
        (coverage_candidates if coverage_candidates is not None else pd.DataFrame()).to_excel(writer, sheet_name="Start_End_Candidates", index=False)
    return output.getvalue()
