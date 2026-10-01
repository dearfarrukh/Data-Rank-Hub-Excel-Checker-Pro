from io import BytesIO

import pandas as pd
from openpyxl import Workbook, load_workbook

from modules.correction_manager import build_force_override, forced_expected, apply_manual_value
from modules.export_manager import build_audit_report_bytes, build_corrected_workbook_bytes
from modules.lifecycle import LifecycleRule
from modules.orientation import detect_orientation, normalize_to_checker
from modules.ranking_engine import build_ranking_matrix, evaluate_missing_cell_risk
from modules.report_builder import build_action_table
from modules.review_manager import apply_resolutions, prepare_action_table
from modules.series_fill import find_safe_fill_candidates, apply_fill_candidates


def test_force_correct_all_earlier_periods():
    rule = build_force_override("Example", "2000", "Production not started yet", "This and all earlier periods")
    assert forced_expected("Example", "1995", [rule])[0] is True
    assert forced_expected("Example", "2000", [rule])[0] is True
    assert forced_expected("Example", "2001", [rule])[0] is False


def test_force_correct_all_later_periods():
    rule = build_force_override("Example", "2005", "Series intentionally ends here", "This and all later periods")
    assert forced_expected("Example", "2004", [rule])[0] is False
    assert forced_expected("Example", "2005", [rule])[0] is True
    assert forced_expected("Example", "2010", [rule])[0] is True


def test_manual_value_changes_only_target_cell():
    df = pd.DataFrame({"Entity": ["A", "B"], "2000": [1, 2], "2001": [None, 3]})
    out, log = apply_manual_value(df, "A", "2001", 9)
    assert out.loc[0, "2001"] == 9
    assert out.loc[1, "2001"] == 3
    assert log["Action"] == "Manual correction"


def test_safe_series_fill_internal_only_and_rounded_for_integer_dataset():
    df = pd.DataFrame({"Entity": ["A"], "2000": [100], "2001": [None], "2002": [121]})
    candidates = find_safe_fill_candidates(df, ["2000", "2001", "2002"], lifecycle_rules=[])
    assert len(candidates) == 1
    # Integer-style datasets preserve 0 decimals, so 110.5 rounds to Python's 110.
    assert candidates.iloc[0]["New Value"] == 110
    out, logs = apply_fill_candidates(df, candidates)
    assert out.loc[0, "2001"] == 110
    assert len(logs) == 1


def test_safe_series_fill_does_not_fill_leading_or_trailing_blanks():
    df = pd.DataFrame({"Entity": ["A"], "2000": [None], "2001": [100], "2002": [120], "2003": [None]})
    candidates = find_safe_fill_candidates(df, ["2000", "2001", "2002", "2003"], lifecycle_rules=[])
    assert candidates.empty


def test_safe_series_fill_skips_protected_transition():
    df = pd.DataFrame({"Entity": ["A"], "1990": [100], "1991": [None], "1992": [200]})
    rules = [LifecycleRule("A", tuple(), None, None, frozenset({1991}), "transition")]
    candidates = find_safe_fill_candidates(df, ["1990", "1991", "1992"], lifecycle_rules=rules)
    assert candidates.empty


def test_boundary_top_n_low_value_is_not_false_risk():
    matrix = pd.DataFrame({"1950": [1000, 900, 800, 700], "1951": [None, 900, 800, 700]})
    ranking = build_ranking_matrix(matrix, 3)
    # row 0's next known value is absent, use reverse arrangement to test boundary from next period
    matrix2 = pd.DataFrame({"1950": [None, 900, 800, 700], "1951": [100, 900, 800, 700]})
    ranking2 = build_ranking_matrix(matrix2, 3)
    risk = evaluate_missing_cell_risk(0, 0, ["1950", "1951"], matrix2, ranking2)
    assert risk["risk"] is False


def test_boundary_top_n_high_value_is_risk():
    matrix = pd.DataFrame({"1950": [None, 900, 800, 700], "1951": [1200, 900, 800, 700]})
    ranking = build_ranking_matrix(matrix, 3)
    risk = evaluate_missing_cell_risk(0, 0, ["1950", "1951"], matrix, ranking)
    assert risk["risk"] is True


def test_resolutions_remove_issue_from_unresolved():
    action = pd.DataFrame([{
        "Severity": "REVIEW", "Entity": "A", "Period": "2000", "Problem": "Test",
        "Current Value": 1, "Previous Value": "", "Previous Rank": "", "Next Value": "",
        "Next Rank": "", "Top-N Cutoff": "", "Why Flagged": "x", "What To Do": "y", "Source": "Core check",
    }])
    prepared = prepare_action_table(action)
    iid = prepared.iloc[0]["Issue ID"]
    unresolved, resolved = apply_resolutions(action, {iid: {"action": "Ignore", "reason": "verified"}})
    assert unresolved.empty
    assert len(resolved) == 1


def test_corrected_export_preserves_other_xlsx_sheet_and_alienart_orientation():
    wb = Workbook()
    ws = wb.active
    ws.title = "AlienArt_Ready"
    ws.append(["Year", "A", "B"])
    ws.append([2000, 1, 2])
    ws.append([2001, None, 3])
    other = wb.create_sheet("Sources")
    other["A1"] = "keep me"
    bio = BytesIO()
    wb.save(bio)
    raw = bio.getvalue()

    original = pd.DataFrame({"Year": [2000, 2001], "A": [1, None], "B": [2, 3]})
    orientation = detect_orientation(original)
    normalized = normalize_to_checker(original, orientation)
    normalized.loc[normalized["Entity"] == "A", "2001"] = 9

    out_bytes = build_corrected_workbook_bytes(raw, "x.xlsx", "AlienArt_Ready", original, normalized, orientation)
    out_wb = load_workbook(BytesIO(out_bytes), data_only=True)
    assert out_wb["Sources"]["A1"].value == "keep me"
    assert out_wb["AlienArt_Ready"]["B3"].value == 9


def test_audit_report_contains_expected_sheets():
    out = build_audit_report_bytes(
        pd.DataFrame(), pd.DataFrame(), [], pd.DataFrame(), pd.DataFrame(), {"Status": "READY"}
    )
    wb = load_workbook(BytesIO(out), read_only=True)
    assert {"Summary", "Unresolved", "Resolved", "Change_Log", "Lifecycle", "Safe_Fill_Preview"}.issubset(set(wb.sheetnames))


def test_add_ranking_context_accepts_arrow_string_columns():
    """Regression: Streamlit Cloud pandas may return Arrow-backed string columns."""
    import pandas as pd
    from types import SimpleNamespace
    from modules.review_manager import add_ranking_context

    try:
        cutoff_col = pd.Series([""], dtype="string[pyarrow]")
    except Exception:
        cutoff_col = pd.Series([""], dtype="string")

    action = pd.DataFrame({
        "Severity": ["MUST CHECK"],
        "Entity": ["Belgium"],
        "Period": ["1961"],
        "Problem": ["Missing"],
        "Top-N Cutoff": cutoff_col,
    })
    working = pd.DataFrame({"Entity": ["Belgium"], "1961": [100.0]})
    ranks = pd.DataFrame({"1961": [9.0]})
    ranking = SimpleNamespace(ranks=ranks, cutoffs={"1961": 75000.0})

    result = add_ranking_context(action, working, ranking)
    assert result.loc[0, "Current Rank"] == 9
    assert result.loc[0, "Top-N Cutoff"] == 75000.0


def test_internal_gap_defaults_to_review_not_must_check():
    from modules.error_checker import CheckerConfig, run_core_checks
    df = pd.DataFrame({"Entity": ["A"], "2000": [100], "2001": [None], "2002": [120]})
    findings = run_core_checks(df, CheckerConfig()).findings_frame()
    gap = findings[findings["Category"].eq("Internal Gap")]
    assert len(gap) == 1
    assert gap.iloc[0]["Severity"] == "REVIEW"


def test_safe_fill_ignores_invalid_text_and_excluded_topn_cells():
    from modules.series_fill import find_safe_fill_candidates
    df = pd.DataFrame({
        "Entity": ["Bad", "Risk", "Safe"],
        "2000": [100, 100, 100],
        "2001": ["BAD_TEXT", None, None],
        "2002": [120, 120, 120],
    })
    out = find_safe_fill_candidates(
        df,
        ["2000", "2001", "2002"],
        lifecycle_rules=[],
        excluded_cells={("Risk", "2001")},
    )
    assert set(zip(out["Entity"], out["Period"])) == {("Safe", "2001")}
    assert float(out.iloc[0]["New Value"]) == 110.0


def test_internal_gap_context_shows_anchor_values_and_ranks():
    from modules.ranking_engine import build_ranking_matrix
    from modules.review_manager import add_ranking_context
    df = pd.DataFrame({"Entity": ["A", "B"], "2000": [100, 90], "2001": [None, 95], "2002": [120, 100]})
    numeric = df[["2000", "2001", "2002"]].apply(pd.to_numeric, errors="coerce")
    ranking = build_ranking_matrix(numeric, 1)
    actions = pd.DataFrame([{\
        "Severity": "REVIEW", "Entity": "A", "Period": "2001", "Problem": "Internal Gap",\
        "Current Value": None, "Previous Value": None, "Previous Rank": None,\
        "Next Value": None, "Next Rank": None, "Top-N Cutoff": None,\
        "Why Flagged": "gap", "What To Do": "review", "Source": "Core check"\
    }])
    out = add_ranking_context(actions, df, ranking)
    assert out.iloc[0]["Previous Value"] == 100.0
    assert out.iloc[0]["Next Value"] == 120.0
    assert out.iloc[0]["Previous Rank"] == 1
    assert out.iloc[0]["Next Rank"] == 1
